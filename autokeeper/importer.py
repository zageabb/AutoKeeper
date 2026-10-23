"""Excel/CSV importer for the uploaded Master Car List layout."""
import csv
import io
import sqlite3

from .core import iso_date, normalise_lifecycle_date, registration, save_vehicle

MAX_IMPORT = 6 * 1024 * 1024


def value(row, *names):
    lookup = {str(k).strip().casefold(): v for k, v in row.items() if k is not None}
    for name in names:
        cell = lookup.get(name.casefold())
        if cell is not None and str(cell).strip():
            return str(cell).strip()
    return ""


def read_import(data, filename):
    if len(data) > MAX_IMPORT:
        raise ValueError("Import files are limited to 6 MB.")
    if filename.lower().endswith(".csv"):
        return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    if filename.lower().endswith(".xlsx"):
        # Import openpyxl only when needed so CSV imports work even if the
        # Ubuntu virtual environment is missing its Excel dependency.
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel import support is not installed. Run '.venv/bin/python -m pip install -r requirements.txt' in the AutoKeeper installation, then restart the service.") from exc
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = book["Master"] if "Master" in book.sheetnames else book.active
            lines = sheet.iter_rows(values_only=True)
            try:
                header = [str(c or "").strip() for c in next(lines)]
            except StopIteration as exc:
                raise ValueError("Spreadsheet is empty") from exc
            return [dict(zip(header, line)) for line in lines]
        finally:
            book.close()
    raise ValueError("Please select a .xlsx or .csv file.")


def import_vehicles(db, rows):
    result = {"created": 0, "updated": 0, "skipped": 0, "errors": []}
    for row_number, row in enumerate(rows, start=2):
        vrm = value(row, "Registration No", "registrationNumber", "Registration", "VRM")
        if not vrm:
            result["skipped"] += 1
            continue
        try:
            vrm = registration(vrm)
            incoming = {
                "registration": vrm,
                "owner": value(row, "Employee Name", "Owner", "Driver"),
                "company": value(row, "Company"),
                "email": value(row, "Email Address", "Email", "Recipient"),
                "make": value(row, "Make", "dvla_make"),
                "model": value(row, "Model", "dvla_Model"),
                "colour": value(row, "Colour"),
            }
            # Empty import cells should preserve existing dates. Nonempty
            # spreadsheet date cells are normalised and validated before saving.
            for field, headings in (
                ("acquisition_date", ("Acquisition Date", "Acquired Date", "acquisition_date", "Date Acquired")),
                ("disposal_date", ("Disposal Date", "Disposed Date", "disposal_date", "Date Disposed")),
            ):
                raw = value(row, *headings)
                if raw:
                    incoming[field] = normalise_lifecycle_date(raw)
            current = db.execute("SELECT * FROM vehicles WHERE registration=?",
                                 (vrm,)).fetchone()
            if current:
                # Blank import cells must not discard existing driver/contact data.
                merged = dict(current)
                merged.update({k: v for k, v in incoming.items() if v})
                saved = save_vehicle(db, merged, current["id"])
                result["updated"] += 1
            else:
                saved = save_vehicle(db, incoming)
                result["created"] += 1
            # Optional imported DVLA Results workbook snapshot. Never replace
            # a record already refreshed live by this application with older
            # spreadsheet data. Snapshot imports are NOT live DVLA checks.
            if not saved["last_checked_at"]:
                mot_raw = value(row, "motExpiryDate", "motExpiry")
                tax_raw = value(row, "taxDueDate", "taxDue")
                mot_date = iso_date(mot_raw)
                tax_date = iso_date(tax_raw)
                mot_status = value(row, "motStatus")
                tax_status = value(row, "taxStatus")
                if any((mot_date, tax_date, mot_status, tax_status)):
                    db.execute("""
                        UPDATE vehicles SET
                          mot_expiry_date=COALESCE(?,mot_expiry_date),
                          tax_due_date=COALESCE(?,tax_due_date),
                          mot_status=COALESCE(NULLIF(?,''),mot_status),
                          tax_status=COALESCE(NULLIF(?,''),tax_status)
                        WHERE id=?
                    """, (mot_date.isoformat() if mot_date else None,
                          tax_date.isoformat() if tax_date else None,
                          mot_status, tax_status, saved["id"]))
                    db.commit()
        except (ValueError, sqlite3.IntegrityError) as exc:
            result["skipped"] += 1
            if len(result["errors"]) < 20:
                result["errors"].append({"row": row_number, "reason": str(exc)})
    return result
